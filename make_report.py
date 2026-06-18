#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Styles ──────────────────────────────────────────────────────────────────
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p

def shade_row(row, hex_color="D9E1F2"):
    for cell in row.cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tcPr.append(shd)

def set_col_width(table, col_idx, width_inches):
    for row in table.rows:
        row.cells[col_idx].width = Inches(width_inches)

# ── Title ────────────────────────────────────────────────────────────────────
doc.add_heading('ADAR Library Design Summary', 0)
doc.add_paragraph(
    'This document summarises the two-part sequence variant library generated '
    'from 9 validated ADAR parent hairpin substrates, together with the '
    'experimental hypotheses underlying each mutation class.'
)

# ════════════════════════════════════════════════════════════════════════════
# TABLE 1 — Library Summary
# ════════════════════════════════════════════════════════════════════════════
heading('Table 1: Library Summary by Parent Substrate', level=1)

doc.add_paragraph(
    'Combined library: 3,076 entries across two sub-libraries (Library 1: 1,800; '
    'Library 2: 1,276). Three parents were excluded as non-simple-hairpin or '
    'edit-site-outside-stem structures (5-HT2C_1, AZIN1_2, GluR_BRG_15mer).'
)

lib1_cols = [
    'Parent', 'Endogenous Substrate', 'Edit Site(s)',
    'Lib 1', 'Lib 2', 'Total',
    'Length Range', 'GC Range'
]

lib1_data = [
    ('5-HT2C_2',        'HTR2C — serotonin receptor 2C pre-mRNA',  'A4, A6, A16, A25, A47',           '200', '200', '400',  '71–75 nt', '30–52%'),
    ('NEIL1',           'NEIL1 — DNA glycosylase pre-mRNA',         'A45',                              '200', '114', '314',  '33–74 nt', '37–68%'),
    ('AZIN1',           'AZIN1 — antizyme inhibitor 1 pre-mRNA',    'A31',                              '200',  '29', '229',  '11–41 nt',  '8–76%'),
    ('BDF2',            '—',                                         'A17',                              '200', '170', '370',  '23–53 nt', '30–67%'),
    ('hGLI1',           'GLI1 — oncogene pre-mRNA',                  'A55',                              '200',  '28', '228',  '30–77 nt', '41–75%'),
    ('PRE_MIRNA142',    'pre-miR-142',                               'A3, A18–20, A23, A37, A47, A52, A59', '200', '200', '400', '66–68 nt', '34–56%'),
    ('GluR-B_QR',       'GRIA2 — glutamate receptor B, Q/R site',   'A6',                               '200', '200', '400',  '42–56 nt', '36–64%'),
    ('GluR-B_RG',       'GRIA2 — glutamate receptor B, R/G site',   'A6',                               '200', '200', '400',  '42–56 nt', '22–52%'),
    ('GluR_BRG_15mer_2','GRIA2 — R/G site 15mer variant',           'A6',                               '200', '135', '335',  '20–34 nt', '23–76%'),
    ('Total',           '',                                           '',                                 '1,800','1,276','3,076', '', ''),
]

t1 = doc.add_table(rows=1, cols=len(lib1_cols))
t1.style = 'Table Grid'
t1.alignment = WD_TABLE_ALIGNMENT.CENTER

# Header
hdr = t1.rows[0]
shade_row(hdr, "2E75B6")
for i, col in enumerate(lib1_cols):
    cell = hdr.cells[i]
    cell.text = col
    run = cell.paragraphs[0].runs[0]
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.size = Pt(10)

# Data rows
for idx, row_data in enumerate(lib1_data):
    row = t1.add_row()
    if idx == len(lib1_data) - 1:  # total row
        shade_row(row, "BDD7EE")
    elif idx % 2 == 0:
        shade_row(row, "DEEAF1")
    for i, val in enumerate(row_data):
        cell = row.cells[i]
        cell.text = val
        run = cell.paragraphs[0].runs[0]
        run.font.size = Pt(9)
        if idx == len(lib1_data) - 1:
            run.bold = True

doc.add_paragraph()

# ════════════════════════════════════════════════════════════════════════════
# Overall Hypothesis
# ════════════════════════════════════════════════════════════════════════════
heading('Overall Hypothesis', level=1)
doc.add_paragraph(
    'ADAR editing efficiency is jointly determined by three kinetic parameters: '
    '(1) kfold — the thermodynamic stability of the dsRNA hairpin conformation; '
    '(2) kbind — the efficiency with which ADAR engages the duplex; and '
    '(3) kflip — the energetic cost of extruding the target adenosine from the '
    'helix for deamination. By systematically varying base-pair identity, stem '
    'length, and the identity of the base opposing the edit site, combined with '
    'DMS-MaPseq structural probing to validate solution conformation, we can '
    'deconvolve the contribution of each parameter to overall editing efficiency '
    'across a diverse set of endogenous ADAR substrates.'
)

# ════════════════════════════════════════════════════════════════════════════
# TABLE 2 — Mutation Hypothesis Table
# ════════════════════════════════════════════════════════════════════════════
heading('Table 2: Mutation Classes and Experimental Hypotheses', level=1)

mut_cols = [
    'Variant Class', 'What Changes', 'Primary Parameter',
    'Hypothesis'
]

mut_data = [
    (
        'Length trimming',
        'Stem shortened inward from both ends',
        'kfold, kbind',
        'A minimum duplex length is required for stable folding and productive '
        'ADAR engagement. Shorter stems reduce kbind by limiting the dsRNA '
        'footprint (~15 bp optimal). DMS-MaPseq confirms which truncations '
        'maintain the native fold.'
    ),
    (
        'Above mutations\n(1–5 bp, cumulative)',
        'bp identity 5′ of edit site changed\nto all-GC or all-AU',
        'kbind, kfold',
        'GC-rich regions above the edit site stabilise the duplex (↑kfold) and '
        'enhance ADAR binding affinity (↑kbind). AU-rich regions destabilise the '
        'duplex, reduce local melting temperature, and lower editing efficiency.'
    ),
    (
        'Below mutations\n(1–5 bp, cumulative)',
        'bp identity 3′ of edit site\n(toward loop) changed\nto all-GC or all-AU',
        'kflip',
        'The bp immediately closing the edit site on the loop side resists base '
        'flipping. GC pairs below the edit site raise the energetic barrier to '
        'kflip and suppress editing; AU pairs facilitate flipping and increase '
        'editing efficiency.'
    ),
    (
        'Across mutations\n(C, G, A)',
        'Base paired with the A edit site\nchanged from U to C, G, or A',
        'kflip, editing rate',
        'The opposing base directly sets the mismatch energy at the edit site. '
        'A·C mismatches are preferred by ADAR over A·U Watson-Crick pairs '
        '(lower kflip barrier). A·G and A·A mismatches test how local instability '
        'drives or inhibits deamination independently of global duplex stability.'
    ),
    (
        'Combined above+below\n(all 5×5 combinations)',
        'Both flanking regions mutated\nsimultaneously, GC or AU',
        'kfold + kflip\n(interaction)',
        'Tests whether the effects of above and below mutations are independent '
        'or cooperative. A stable duplex above with a weak bp below may optimally '
        'balance kbind and kflip; fully GC or fully AU surroundings reveal the '
        'upper and lower bounds of editing efficiency.'
    ),
    (
        'DMS-MaPseq\nstructural validation',
        'No sequence change — chemical\nprobing of all variants',
        'kfold (structural)',
        'Confirms that designed mutations produce the intended structural changes '
        'in solution. Unpaired A/C reactivity maps validate the assumed dot-bracket '
        'topology and reveal unexpected misfolding or alternative structures that '
        'could confound kinetic interpretation.'
    ),
]

t2 = doc.add_table(rows=1, cols=len(mut_cols))
t2.style = 'Table Grid'
t2.alignment = WD_TABLE_ALIGNMENT.CENTER

hdr2 = t2.rows[0]
shade_row(hdr2, "2E75B6")
for i, col in enumerate(mut_cols):
    cell = hdr2.cells[i]
    cell.text = col
    run = cell.paragraphs[0].runs[0]
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.size = Pt(10)

for idx, row_data in enumerate(mut_data):
    row = t2.add_row()
    if idx % 2 == 0:
        shade_row(row, "DEEAF1")
    for i, val in enumerate(row_data):
        cell = row.cells[i]
        cell.text = val
        run = cell.paragraphs[0].runs[0]
        run.font.size = Pt(9)
        if i == 0:
            run.bold = True

# ── Save ────────────────────────────────────────────────────────────────────
out = '/home/user/RNA/ADAR_Library_Summary.docx'
doc.save(out)
print(f'saved {out}')
